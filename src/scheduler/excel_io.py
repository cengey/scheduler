"""Excel import and export utilities for distributions and Sandbox schedules."""

import io
from datetime import datetime, time, date, timedelta
from typing import Any, Dict, List, Optional, Tuple, Union
import openpyxl

DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def parse_distribution_workbook(
    file_source: Union[str, bytes, io.BytesIO],
    class_type: str,
) -> Dict[str, Any]:
    """Parse the last tab of a distribution workbook.

    Finds the 'CET Time' row and extracts the 48 half-hour slots across Monday-Sunday.
    Returns:
        {
            "class_type": str,
            "sheet_name": str,
            "slots": { (day_index_0_to_6, "HH:MM"): count },
            "total_count": int,
            "days_totals": [int, ... 7 days]
        }
    """
    if isinstance(file_source, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_source), data_only=True)
    elif isinstance(file_source, io.BytesIO):
        wb = openpyxl.load_workbook(file_source, data_only=True)
    else:
        wb = openpyxl.load_workbook(file_source, data_only=True)

    ws = wb.worksheets[-1]  # Always use the last tab as required
    sheet_name = ws.title

    # Find the row containing 'CET Time'
    header_row_idx = None
    time_col_idx = None
    day_col_indices: List[Tuple[int, int]] = []  # (day_idx 0..6, col_idx 1-based)

    for r_idx in range(1, min(15, ws.max_row + 1)):
        for c_idx in range(1, min(15, ws.max_column + 1)):
            val = str(ws.cell(r_idx, c_idx).value or "").strip().lower()
            if "cet" in val and "time" in val:
                header_row_idx = r_idx
                time_col_idx = c_idx
                break
        if header_row_idx:
            break

    if not header_row_idx or not time_col_idx:
        # Fallback to standard row 3, col 1
        header_row_idx = 3
        time_col_idx = 1

    # Map the 7 day columns following time_col_idx
    for c_idx in range(time_col_idx + 1, time_col_idx + 15):
        val = str(ws.cell(header_row_idx, c_idx).value or "").strip()
        for d_idx, d_name in enumerate(DAY_NAMES):
            if d_name.lower() in val.lower():
                day_col_indices.append((d_idx, c_idx))
                break

    # If day headers were not named or parsed, default to next 7 consecutive columns
    if len(day_col_indices) < 7:
        day_col_indices = [(i, time_col_idx + 1 + i) for i in range(7)]

    slots: Dict[Tuple[int, str], int] = {}
    days_totals = [0] * 7
    total_count = 0

    # Scan the 48 time slots (00:00 to 23:30)
    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        cell_time = ws.cell(r_idx, time_col_idx).value
        slot_str = None
        if isinstance(cell_time, time):
            slot_str = cell_time.strftime("%H:%M")
        elif isinstance(cell_time, datetime):
            slot_str = cell_time.strftime("%H:%M")
        elif isinstance(cell_time, str) and ":" in cell_time:
            # Handle text strings like '14:00' or '14:00:00'
            parts = cell_time.strip().split(":")
            if len(parts) >= 2:
                slot_str = f"{int(parts[0]):02d}:{int(parts[1]):02d}"

        if not slot_str:
            continue

        for day_idx, col_idx in day_col_indices:
            raw_val = ws.cell(r_idx, col_idx).value
            count = 0
            if isinstance(raw_val, (int, float)):
                count = int(raw_val)
            elif isinstance(raw_val, str) and raw_val.strip().isdigit():
                count = int(raw_val.strip())

            if count > 0:
                slots[(day_idx, slot_str)] = count
                days_totals[day_idx] += count
                total_count += count

    return {
        "class_type": class_type,
        "sheet_name": sheet_name,
        "slots": slots,
        "total_count": total_count,
        "days_totals": days_totals,
    }


def parse_schedule_file(
    file_source: Union[str, bytes, io.BytesIO],
) -> List[Dict[str, Any]]:
    """Parse an existing schedule workbook (matching the Sandbox schedule format).

    Columns: Account, Name, Time, End Time, VTC, Class Type, Hidden
    """
    if isinstance(file_source, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_source), data_only=True)
    elif isinstance(file_source, io.BytesIO):
        wb = openpyxl.load_workbook(file_source, data_only=True)
    else:
        wb = openpyxl.load_workbook(file_source, data_only=True)

    ws = wb.worksheets[-1]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        return []

    # Map column headers
    header = [str(c or "").strip().lower() for c in rows[0]]
    col_map = {
        "account": 0,
        "name": 1,
        "time": 2,
        "end time": 3,
        "vtc": 4,
        "class type": 5,
        "hidden": 6,
    }
    for idx, h in enumerate(header):
        for k in col_map:
            if k in h:
                col_map[k] = idx

    classes: List[Dict[str, Any]] = []
    for r in rows[1:]:
        if not r or r[col_map["account"]] is None:
            continue

        def _fmt_dt(val: Any) -> str:
            if isinstance(val, (datetime, date)):
                if isinstance(val, date) and not isinstance(val, datetime):
                    val = datetime.combine(val, time.min)
                return val.strftime("%Y-%m-%d %H:%M:%S")
            return str(val or "").strip()

        account = str(r[col_map["account"]]).strip()
        name = str(r[col_map["name"]]).strip()
        start_t = _fmt_dt(r[col_map["time"]])
        end_t = _fmt_dt(r[col_map["end time"]])
        vtc = str(r[col_map["vtc"]]).strip() if len(r) > col_map["vtc"] and r[col_map["vtc"]] else ""
        ctype = str(r[col_map["class type"]]).strip() if len(r) > col_map["class type"] and r[col_map["class type"]] else "Adult"
        hidden = str(r[col_map["hidden"]]).strip() if len(r) > col_map["hidden"] and r[col_map["hidden"]] else "No"

        classes.append({
            "account": account,
            "name": name,
            "time": start_t,
            "end_time": end_t,
            "vtc": vtc,
            "class_type": ctype,
            "hidden": hidden,
        })

    return classes


def export_schedule_to_excel(
    classes: List[Dict[str, Any]],
    output_path_or_buffer: Optional[Union[str, io.BytesIO]] = None,
) -> bytes:
    """Export schedule classes to an Excel file with the exact Sandbox Schedule structure.

    Columns: Account, Name, Time, End Time, VTC, Class Type, Hidden
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Classes"

    headers = ["Account", "Name", "Time", "End Time", "VTC", "Class Type", "Hidden"]
    ws.append(headers)

    # Sort classes chronologically by Time, then Teacher Name
    sorted_classes = sorted(
        classes,
        key=lambda c: (str(c.get("time", "")), str(c.get("name", "")), str(c.get("class_type", "")))
    )

    for item in sorted_classes:
        ws.append([
            item.get("account", ""),
            item.get("name", ""),
            item.get("time", ""),
            item.get("end_time", ""),
            item.get("vtc", ""),
            item.get("class_type", ""),
            item.get("hidden", "No"),
        ])

    # Auto-fit column widths slightly for neatness
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    if isinstance(output_path_or_buffer, str):
        wb.save(output_path_or_buffer)
        with open(output_path_or_buffer, "rb") as f:
            return f.read()
    elif isinstance(output_path_or_buffer, io.BytesIO):
        wb.save(output_path_or_buffer)
        return output_path_or_buffer.getvalue()
    else:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def parse_hub_teacher_list(
    file_source: Union[str, bytes, io.BytesIO],
) -> List[Dict[str, Any]]:
    """Parse Hub Teacher List.xlsx to extract master teacher profiles.

    Columns expected: Name, VTC, Status, Platform Username, Email, etc.
    """
    if isinstance(file_source, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_source), data_only=True)
    elif isinstance(file_source, io.BytesIO):
        wb = openpyxl.load_workbook(file_source, data_only=True)
    else:
        wb = openpyxl.load_workbook(file_source, data_only=True)

    ws = wb['Teachers'] if 'Teachers' in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    if not rows:
        return []

    header = [str(c or "").strip().lower() for c in rows[0]]
    col_map = {
        "name": 0,
        "vtc": 1,
        "status": 2,
        "username": 3,
        "email": 5,
        "shift": None,
        "rotation": None,
        "day_off": None,
        "weekend": None,
        "teens": None,
        "speak_plus": None,
    }

    for idx, h in enumerate(header):
        if "user" in h or "account" in h or "login" in h:
            col_map["username"] = idx
        elif "vtc" in h:
            col_map["vtc"] = idx
        elif "status" in h:
            col_map["status"] = idx
        elif "name" in h and "user" not in h:
            col_map["name"] = idx
        elif "email" in h:
            col_map["email"] = idx
        elif "shift" in h:
            col_map["shift"] = idx
        elif "rotation" in h or "group" in h:
            col_map["rotation"] = idx
        elif "off" in h:
            col_map["day_off"] = idx
        elif "teen" in h:
            col_map["teens"] = idx
        elif "speak" in h:
            col_map["speak_plus"] = idx

    teachers: List[Dict[str, Any]] = []
    for r in rows[1:]:
        if not r or len(r) <= (col_map["username"] or 3):
            continue
        username = str(r[col_map["username"]] or "").strip()
        if not username:
            continue

        name = str(r[col_map["name"]] or username).strip()
        vtc = str(r[col_map["vtc"]] or "").strip() if col_map["vtc"] is not None and len(r) > col_map["vtc"] else ""
        raw_status = str(r[col_map["status"]] or "Active").strip() if col_map["status"] is not None and len(r) > col_map["status"] else "Active"
        email = str(r[col_map["email"]] or "").strip() if col_map["email"] is not None and len(r) > col_map["email"] else ""

        status = "resigned" if raw_status.lower() in ("inactive", "resigned") else "active"
        if "to be assigned" in name.lower() or name.lower() == "cengey g":
            status = "placeholder"

        item = {
            "account": username,
            "name": name,
            "vtc": vtc,
            "email": email,
            "status": status,
        }

        if col_map["shift"] is not None and len(r) > col_map["shift"] and r[col_map["shift"]]:
            item["shift_id"] = str(r[col_map["shift"]]).strip().lower()
        if col_map["rotation"] is not None and len(r) > col_map["rotation"] and r[col_map["rotation"]]:
            item["rotation_id"] = str(r[col_map["rotation"]]).strip()
        if col_map["day_off"] is not None and len(r) > col_map["day_off"] and r[col_map["day_off"]]:
            item["fixed_day_off"] = str(r[col_map["day_off"]]).strip()
        if col_map["teens"] is not None and len(r) > col_map["teens"] and r[col_map["teens"]]:
            t_val = str(r[col_map["teens"]]).strip().lower()
            item["teens_qualified"] = t_val in ("1", "true", "yes", "y")
        if col_map["speak_plus"] is not None and len(r) > col_map["speak_plus"] and r[col_map["speak_plus"]]:
            sp_val = str(r[col_map["speak_plus"]]).strip().lower()
            item["speak_plus_qualified"] = sp_val in ("1", "true", "yes", "y")

        teachers.append(item)

    return teachers


def export_teachers_to_excel(teachers: List[Dict[str, Any]]) -> bytes:
    """Export the master teacher roster to an Excel workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Teachers"

    headers = [
        "Account / Username", "Full Name", "VTC", "Email", "Status",
        "Shift", "Rotation Group", "Fixed Day Off", "Weekend Group",
        "Teens Qualified", "Speak+ Qualified"
    ]
    ws.append(headers)

    for t in teachers:
        ws.append([
            t.get("account", ""),
            t.get("name", ""),
            t.get("vtc", ""),
            t.get("email", ""),
            t.get("status", "active"),
            t.get("shift_id", "day"),
            t.get("rotation_id", "rot-1"),
            t.get("fixed_day_off", "Thursday"),
            t.get("weekend_group", "A"),
            "Yes" if t.get("teens_qualified") else "No",
            "Yes" if t.get("speak_plus_qualified") else "No",
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
